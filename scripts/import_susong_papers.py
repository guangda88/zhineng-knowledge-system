#!/usr/bin/env python3
"""
苏颂研究论文批量导入脚本

从百度云(alist)下载论文 → 解析文本 → 导入 documents 表
支持: PDF (PyPDF2), DOCX (python-docx)
不支持: CAJ (知网格式，需人工转换)

用法:
    python scripts/import_susong_papers.py                    # 导入所有论文
    python scripts/import_susong_papers.py --category 天文科技  # 只导入某个分类
    python scripts/import_susong_papers.py --dry-run          # 预览不执行
    python scripts/import_susong_papers.py --skip-download     # 跳过下载，只解析本地文件
"""

import argparse
import asyncio
import json
import os
import sys
import time
import traceback
from pathlib import Path

import asyncpg

# Add backend to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ALIST_URL = "http://localhost:4255/api/fs/list"
ALIST_DL_URL = "http://localhost:4255/api/fs/get"
CLOUD_BASE = "/百度云9080/ZNQG/注意保管的资料/苏颂研究/参考文档/论文"
LOCAL_BASE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "susong_import")
DB_DSN = os.getenv("DATABASE_URL")

DOWNLOAD_DELAY = 1.5  # seconds between downloads to avoid throttling

CATEGORY_MAP = {
    "天文科技": {"db_category": "科学", "tags": ["苏颂研究", "天文科技"]},
    "政治思想": {"db_category": "哲学", "tags": ["苏颂研究", "政治思想"]},
    "外交关系": {"db_category": "哲学", "tags": ["苏颂研究", "外交关系"]},
    "生平年考": {"db_category": "哲学", "tags": ["苏颂研究", "生平年考"]},
    "教育思想": {"db_category": "儒家", "tags": ["苏颂研究", "教育思想"]},
    "理学道学": {"db_category": "儒家", "tags": ["苏颂研究", "理学道学"]},
    "史学文学": {"db_category": "儒家", "tags": ["苏颂研究", "史学文学"]},
    "综合评价": {"db_category": "哲学", "tags": ["苏颂研究", "综合评价"]},
    "中医中药": {"db_category": "中医", "tags": ["苏颂研究", "中医中药"]},
    "新仪象法要": {"db_category": "科学", "tags": ["苏颂研究", "新仪象法要"]},
}

SKIP_EXTENSIONS = {".caj"}  # CAJ files cannot be parsed


def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


async def alist_list(path, per_page=200):
    """List directory via alist API."""
    import httpx
    async with httpx.AsyncClient(timeout=30) as client:
        r = await client.post(ALIST_URL, json={"path": path, "password": "", "page": 1, "per_page": per_page})
        data = r.json()
        if data.get("code") != 200:
            raise RuntimeError(f"alist error: {data.get('message', 'unknown')}")
        return data["data"]["content"]


async def alist_download(client, remote_path, local_path):
    """Download a file from alist."""
    r = await client.post(ALIST_DL_URL, json={"path": remote_path, "password": ""})
    data = r.json()
    if data.get("code") != 200:
        raise RuntimeError(f"alist download error: {data.get('message', 'unknown')}")

    raw_url = data["data"]["raw_url"]
    if not raw_url:
        raise RuntimeError("No raw_url returned")

    resp = await client.get(raw_url, follow_redirects=True, timeout=120)
    resp.raise_for_status()

    os.makedirs(os.path.dirname(local_path), exist_ok=True)
    with open(local_path, "wb") as f:
        f.write(resp.content)

    return len(resp.content)


def parse_pdf(file_path):
    """Extract text from PDF using PyPDF2."""
    import PyPDF2
    pages = []
    with open(file_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            text = page.extract_text()
            if text and text.strip():
                pages.append(text.strip())
    return "\n\n".join(pages)


def parse_docx(file_path):
    """Extract text from DOCX using python-docx."""
    from docx import Document
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                tables_text.append(" | ".join(row_text))

    content = "\n".join(paragraphs)
    if tables_text:
        content += "\n\n" + "\n".join(tables_text)
    return content


def parse_txt(file_path):
    """Read text file with encoding detection."""
    for encoding in ["utf-8", "gbk", "gb18030", "utf-16"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(file_path, "r", encoding="latin-1", errors="ignore") as f:
        return f.read()


def parse_doc(file_path):
    """Extract text from .doc (old format) using catdoc or antiword."""
    import subprocess
    for cmd in ["catdoc", "antiword"]:
        try:
            result = subprocess.run(
                [cmd, file_path], capture_output=True, text=True, timeout=30
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout.strip()
        except Exception:
            continue
    return None


def parse_file(file_path):
    """Parse file and return text content."""
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        text = parse_pdf(file_path)
    elif ext == ".docx":
        text = parse_docx(file_path)
    elif ext == ".doc":
        try:
            text = parse_docx(file_path)
        except Exception:
            text = None
        if not text or len(text.strip()) < 10:
            text = parse_doc(file_path)
    elif ext in (".txt", ".md", ".markdown"):
        text = parse_txt(file_path)
    else:
        return None
    if text:
        text = text.replace("\x00", "")
    return text


def title_from_filename(filename):
    """Extract clean title from filename."""
    name = filename
    for ext in [".pdf", ".docx", ".doc", ".txt", ".md", ".caj"]:
        name = name.replace(ext, "")
    name = name.strip()
    if name.startswith("《") and name.endswith("》"):
        name = name[1:-1]
    return name


async def check_existing(pool, title):
    """Check if a document with this title already exists."""
    row = await pool.fetchrow(
        "SELECT id FROM documents WHERE title = $1 AND '苏颂研究' = ANY(tags) LIMIT 1",
        title
    )
    return row is not None


async def insert_document(pool, title, content, category, tags):
    """Insert document into database."""
    doc_id = await pool.fetchval(
        """INSERT INTO documents (title, content, category, tags)
           VALUES ($1, $2, $3, $4)
           ON CONFLICT DO NOTHING
           RETURNING id""",
        title, content, category, tags
    )
    return doc_id


async def import_category(pool, category_name, dry_run=False, skip_download=False):
    """Import all papers from a cloud category."""
    if category_name not in CATEGORY_MAP:
        log(f"  Unknown category: {category_name}, skipping")
        return {"total": 0, "success": 0, "skipped": 0, "failed": 0, "errors": []}

    config = CATEGORY_MAP[category_name]
    db_category = config["db_category"]
    base_tags = config["tags"]
    cloud_dir = f"{CLOUD_BASE}/{category_name}"
    local_dir = os.path.join(LOCAL_BASE, category_name)

    log(f"=== Importing [{category_name}] -> DB category: {db_category} ===")

    try:
        items = await alist_list(cloud_dir)
    except Exception as e:
        log(f"  Failed to list cloud directory: {e}")
        return {"total": 0, "success": 0, "skipped": 0, "failed": 1, "errors": [str(e)]}

    log(f"  Found {len(items)} files in cloud")

    stats = {"total": len(items), "success": 0, "skipped": 0, "failed": 0, "errors": []}
    import httpx

    async with httpx.AsyncClient(timeout=120) as client:
        for i, item in enumerate(items):
            name = item.get("name", "")
            ext = Path(name).suffix.lower()
            title = title_from_filename(name)

            # Skip unsupported formats
            if ext in SKIP_EXTENSIONS:
                log(f"  [{i+1}/{len(items)}] SKIP (CAJ): {name}")
                stats["skipped"] += 1
                continue

            # Check if already exists
            try:
                exists = await check_existing(pool, title)
                if exists:
                    log(f"  [{i+1}/{len(items)}] SKIP (exists): {name}")
                    stats["skipped"] += 1
                    continue
            except Exception:
                pass

            local_path = os.path.join(local_dir, name)

            # Download if needed
            if not skip_download:
                if not os.path.exists(local_path):
                    remote_path = f"{cloud_dir}/{name}"
                    try:
                        size = await alist_download(client, remote_path, local_path)
                        log(f"  [{i+1}/{len(items)}] Downloaded: {name} ({size/1024:.0f} KB)")
                        await asyncio.sleep(DOWNLOAD_DELAY)
                    except Exception as e:
                        log(f"  [{i+1}/{len(items)}] DOWNLOAD FAILED: {name} - {e}")
                        stats["failed"] += 1
                        stats["errors"].append(f"download:{name}:{e}")
                        await asyncio.sleep(DOWNLOAD_DELAY)
                        continue
                else:
                    log(f"  [{i+1}/{len(items)}] Using cached: {name}")

            # Parse file
            if not os.path.exists(local_path):
                log(f"  [{i+1}/{len(items)}] SKIP (no local file): {name}")
                stats["skipped"] += 1
                continue

            try:
                content = parse_file(local_path)
                if not content or len(content.strip()) < 50:
                    log(f"  [{i+1}/{len(items)}] SKIP (empty/too short): {name} ({len(content.strip()) if content else 0} chars)")
                    stats["skipped"] += 1
                    continue
            except Exception as e:
                log(f"  [{i+1}/{len(items)}] PARSE FAILED: {name} - {e}")
                stats["failed"] += 1
                stats["errors"].append(f"parse:{name}:{e}")
                continue

            if dry_run:
                log(f"  [{i+1}/{len(items)}] DRY RUN: would import '{title}' ({len(content)} chars)")
                stats["success"] += 1
                continue

            # Insert into DB
            try:
                doc_id = await insert_document(pool, title, content, db_category, base_tags)
                if doc_id:
                    log(f"  [{i+1}/{len(items)}] IMPORTED: '{title}' -> ID {doc_id} ({len(content)} chars)")
                    stats["success"] += 1
                else:
                    log(f"  [{i+1}/{len(items)}] SKIP (duplicate DB): {name}")
                    stats["skipped"] += 1
            except Exception as e:
                log(f"  [{i+1}/{len(items)}] DB INSERT FAILED: {name} - {e}")
                stats["failed"] += 1
                stats["errors"].append(f"db:{name}:{e}")

    log(f"  Category [{category_name}] done: {stats['success']} imported, {stats['skipped']} skipped, {stats['failed']} failed")
    return stats


async def import_xinyi(dry_run=False, skip_download=False):
    """Import 新仪象法要 files from 苏颂全集 root."""
    pool = await asyncpg.create_pool(DB_DSN, min_size=1, max_size=3)
    try:
        cloud_dir = "/百度云9080/ZNQG/注意保管的资料/苏颂研究/苏颂全集"
        local_dir = os.path.join(LOCAL_BASE, "新仪象法要")

        log("=== Importing [新仪象法要] from 苏颂全集 ===")

        try:
            items = await alist_list(cloud_dir, per_page=200)
        except Exception as e:
            log(f"  Failed to list: {e}")
            return

        xinyi_files = [
            item for item in items
            if not item.get("is_dir") and any(
                kw in item.get("name", "")
                for kw in ["新仪象法要", "仪象法纂"]
            )
        ]

        log(f"  Found {len(xinyi_files)} 新仪象法要/仪象法纂 files")

        import httpx
        stats = {"total": len(xinyi_files), "success": 0, "skipped": 0, "failed": 0, "errors": []}

        async with httpx.AsyncClient(timeout=120) as client:
            for i, item in enumerate(xinyi_files):
                name = item.get("name", "")
                ext = Path(name).suffix.lower()
                title = title_from_filename(name)

                if ext in SKIP_EXTENSIONS:
                    stats["skipped"] += 1
                    continue

                exists = await check_existing(pool, title)
                if exists:
                    log(f"  [{i+1}/{len(xinyi_files)}] SKIP (exists): {name}")
                    stats["skipped"] += 1
                    continue

                local_path = os.path.join(local_dir, name)

                if not skip_download and not os.path.exists(local_path):
                    remote_path = f"{cloud_dir}/{name}"
                    try:
                        size = await alist_download(client, remote_path, local_path)
                        log(f"  [{i+1}/{len(xinyi_files)}] Downloaded: {name} ({size/1024:.0f} KB)")
                        await asyncio.sleep(DOWNLOAD_DELAY)
                    except Exception as e:
                        log(f"  [{i+1}/{len(xinyi_files)}] DOWNLOAD FAILED: {name} - {e}")
                        stats["failed"] += 1
                        await asyncio.sleep(DOWNLOAD_DELAY)
                        continue

                if not os.path.exists(local_path):
                    stats["skipped"] += 1
                    continue

                try:
                    content = parse_file(local_path)
                    if not content or len(content.strip()) < 50:
                        stats["skipped"] += 1
                        continue
                except Exception as e:
                    log(f"  [{i+1}/{len(xinyi_files)}] PARSE FAILED: {name} - {e}")
                    stats["failed"] += 1
                    continue

                if dry_run:
                    log(f"  DRY RUN: would import '{title}' ({len(content)} chars)")
                    stats["success"] += 1
                    continue

                try:
                    doc_id = await insert_document(pool, title, content, "科学", ["苏颂研究", "新仪象法要"])
                    if doc_id:
                        log(f"  IMPORTED: '{title}' -> ID {doc_id}")
                        stats["success"] += 1
                    else:
                        stats["skipped"] += 1
                except Exception as e:
                    log(f"  DB FAILED: {name} - {e}")
                    stats["failed"] += 1

        log(f"  新仪象法要 done: {stats['success']} imported, {stats['skipped']} skipped, {stats['failed']} failed")
    finally:
        await pool.close()


async def main():
    parser = argparse.ArgumentParser(description="苏颂研究论文批量导入")
    parser.add_argument("--category", "-c", help="只导入某个分类 (天文科技/政治思想/外交关系/生平年考/教育思想/理学道学/史学文学/综合评价/中医中药)")
    parser.add_argument("--dry-run", "-n", action="store_true", help="预览模式，不实际导入")
    parser.add_argument("--skip-download", action="store_true", help="跳过下载，只处理本地已有文件")
    parser.add_argument("--xinyi", action="store_true", help="导入新仪象法要/仪象法纂文件")
    args = parser.parse_args()

    log("苏颂研究论文批量导入工具")
    log(f"Dry run: {args.dry_run}")

    if args.xinyi:
        await import_xinyi(dry_run=args.dry_run, skip_download=args.skip_download)
        return

    pool = await asyncpg.create_pool(DB_DSN, min_size=1, max_size=3)
    try:
        if args.category:
            categories = [args.category]
        else:
            categories = list(CATEGORY_MAP.keys())

        total_stats = {"total": 0, "success": 0, "skipped": 0, "failed": 0, "errors": []}

        for cat in categories:
            stats = await import_category(pool, cat, dry_run=args.dry_run, skip_download=args.skip_download)
            for k in ["total", "success", "skipped", "failed"]:
                total_stats[k] += stats.get(k, 0)
            total_stats["errors"].extend(stats.get("errors", []))
            log("")

        log("=" * 60)
        log(f"TOTAL: {total_stats['success']} imported, {total_stats['skipped']} skipped, {total_stats['failed']} failed")
        if total_stats["errors"]:
            log(f"Errors ({len(total_stats['errors'])}):")
            for e in total_stats["errors"][:20]:
                log(f"  - {e}")
    finally:
        await pool.close()


if __name__ == "__main__":
    asyncio.run(main())
