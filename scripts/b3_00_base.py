#!/usr/bin/env python3
"""基础框架：封面+前言+目录"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os, pickle

OUTDIR = "/home/ai/下载/未来医学相关/师资培训教材"
OUTPATH = os.path.join(OUTDIR, "第三册_医学气功与导引功法_师资培训教材.docx")
os.makedirs(OUTDIR, exist_ok=True)
STATE = "/tmp/b3_progress.pkl"

doc = Document()
style = doc.styles['Normal']
style.font.name = '宋体'; style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.line_spacing = 1.5; style.paragraph_format.space_after = Pt(4)
for sec in doc.sections:
    sec.top_margin = Cm(2.54); sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17); sec.right_margin = Cm(3.17)

total_chars = 0

def H(text, level=1):
    sz = {0:22,1:18,2:16,3:14,4:13}[level]
    fn = {0:'黑体',1:'黑体',2:'黑体',3:'楷体',4:'楷体'}[level]
    p = doc.add_paragraph()
    if level <= 2: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text); run.bold = True
    run.font.size = Pt(sz); run.font.name = fn
    run.element.rPr.rFonts.set(qn('w:eastAsia'), fn)
    p.paragraph_format.space_before = Pt({0:30,1:24,2:18,3:12,4:8}[level])
    p.paragraph_format.space_after = Pt({0:20,1:14,2:10,3:8,4:6}[level])
    if level == 2: p.paragraph_format.page_break_before = True
    global total_chars; total_chars += len(text)

def P(text):
    global total_chars; total_chars += len(text)
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12); run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5

def PB():
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)

def TBL(headers, rows, cw=None):
    global total_chars
    for r in rows: total_chars += sum(len(str(v)) for v in r)
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c = table.rows[0].cells[i]; c.text = ''
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.bold = True; r.font.size = Pt(11)
        r.font.name = '宋体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="D9E2F3"/>'))
    for ri, rd in enumerate(rows):
        for ci, v in enumerate(rd):
            c = table.rows[ri+1].cells[ci]; c.text = ''
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(v)); r.font.size = Pt(11)
            r.font.name = '宋体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if cw:
        for row in table.rows:
            for i,w in enumerate(cw):
                if i < len(row.cells): row.cells[i].width = Cm(w)
    PB()

# ===== 封面 =====
for _ in range(6): PB()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('智能健康与康养·师资培训教材'); r.bold = True; r.font.size = Pt(16)
r.font.name = '宋体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
r.font.color.rgb = RGBColor(0x44,0x72,0xC4)
PB()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('第 三 册'); r.bold = True; r.font.size = Pt(36)
r.font.name = '黑体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
r.font.color.rgb = RGBColor(0x1F,0x49,0x7D)
PB()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('《医学气功与导引功法》'); r.bold = True; r.font.size = Pt(28)
r.font.name = '黑体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
r.font.color.rgb = RGBColor(0x1F,0x49,0x7D)
PB()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('师 资 培 训 教 材'); r.bold = True; r.font.size = Pt(24)
r.font.name = '黑体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
for _ in range(4): PB()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('定位：功法带教老师、康养教练专用'); r.font.size = Pt(14)
r.font.name = '楷体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('总课时：80课时'); r.font.size = Pt(14)
r.font.name = '楷体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
doc.add_page_break()

# ===== 前言 =====
H('前  言', 1)
P('气功是中华民族传统文化的瑰宝，是中医学的重要组成部分。医学气功作为气功的一个重要分支，以中医理论为指导，通过调身、调息、调心的综合锻炼，达到强身健体、防病治病、延年益寿的目的。数千年来，医学气功在维护中华民族的健康中发挥了不可替代的作用，积累了丰富的理论体系和实践经验。从《黄帝内经》中"恬淡虚无，真气从之，精神内守，病安从来"的养生智慧，到华佗创编五禽戏的导引实践，从孙思邈的养生导引法到现代八段锦、太极拳的广泛推广，医学气功始终是中华民族健康文化的重要载体。')
P('随着健康中国战略的深入推进，人民群众对健康的需求日益增长，对非药物健康干预方法的关注度持续提高。医学气功作为一种绿色、安全、有效的健康干预手段，正在被越来越多的人所接受和实践。特别是在慢性病防治、亚健康调理、心理健康促进等领域，医学气功展现出了独特的优势和广阔的应用前景。现代科学研究也不断证实，规律的气功锻炼可以改善心血管功能、调节自主神经系统、增强免疫功能、缓解压力和焦虑，对多种慢性疾病具有辅助治疗作用。')
P('导引功法是医学气功的重要实践形式，它融合了呼吸运动、肢体运动和意念活动，是一种身心兼修的锻炼方法。从马王堆汉墓出土的导引图到华佗创编的五禽戏，从孙思邈的养生导引法到现代的八段锦、太极拳，导引功法历经数千年的传承与发展，形成了丰富完整的理论体系和实践方法。导引功法动作舒缓、简单易学、安全可靠，适合不同年龄、不同体质、不同健康状况的人群练习，是实施全民健身计划和健康中国行动的重要抓手。')
P('本教材作为《智能健康与康养·师资培训教材》系列的第三册，定位为功法带教老师和康养教练的专用教材。全书共分十章，涵盖医学气功基础理论、呼吸训练教学法、站桩功法体系教学、八段锦标准化教学、简化太极拳教学、坐式与卧式导引、颈肩腰专项康复功法、睡眠减压安神专项功法、不同人群功法适配教学以及团体带教与活动组织等内容。总课时80课时，每章8课时，理论与实操并重。')
P('本教材的编写遵循"理论够用、突出实操、安全第一、易于教学"的原则，力求将传统功法的精髓与现代教学理念相结合，为功法带教老师提供系统完整的教学依据。每一章节均包含理论基础、操作要领、教学要点、常见问题、安全规范等内容，并配有教学表格供参考使用。同时，教材特别强调了安全意识和风险防控，确保教学过程规范、安全、有效。')
P('本教材适用于各类健康管理机构、康养中心、社区党群服务中心、工会职工之家、退役军人服务站等场所的功法带教老师和康养教练培训使用。希望本教材能够帮助广大功法带教老师掌握规范的教学方法，提高教学质量，为推动医学气功事业的健康发展、增进人民群众健康福祉贡献力量。')
doc.add_page_break()

# ===== 目录 =====
H('目  录', 1)
toc = [
('第1章','医学气功基础理论','8课时'),('  1.1','医学气功定义、历史、现代价值',''),
('  1.2','调身、调息、调心三要素',''),('  1.3','气血、经络、精气神基础',''),
('  1.4','气功科学机理：神经、内分泌、免疫',''),('  1.5','安全原则、禁忌、适宜人群',''),
('第2章','呼吸训练教学法','8课时'),('  2.1','自然呼吸、腹式呼吸、逆腹式呼吸',''),
('  2.2','吐纳、闭气、调息规范',''),('  2.3','呼吸与情绪、睡眠、疼痛调节',''),
('  2.4','站、坐、卧三姿呼吸练习',''),('  2.5','常见问题纠正与安全把控',''),
('第3章','站桩功法体系教学','8课时'),('  3.1','无极桩：身形、放松、入静',''),
('  3.2','混元桩：要领、呼吸、意念',''),('  3.3','桩功训练步骤、时长、频次',''),
('  3.4','肩紧、腰累、头晕、走神纠正',''),('  3.5','桩功功效与日常保健应用',''),
('第4章','八段锦标准化教学','8课时'),('  4.1','八段锦源流、功效、适应人群',''),
('  4.2','一式至八式动作分解与口令',''),('  4.3','呼吸配合、意念引导、发力要点',''),
('  4.4','动作纠错、常见错误纠正',''),('  4.5','团体带练流程与音乐配合',''),
('第5章','简化太极拳教学','8课时'),('  5.1','太极基础：松、静、柔、缓',''),
('  5.2','起势、收势、核心动作',''),('  5.3','云手、揽雀尾、单鞭等基础式',''),
('  5.4','节奏、呼吸、意念训练',''),('  5.5','老年、慢病、体弱适配版',''),
('第6章','坐式导引与卧式导引','8课时'),('  6.1','坐式放松功：办公、居家、养老适用',''),
('  6.2','坐式八段锦、关节活动操',''),('  6.3','卧式放松、助眠导引',''),
('  6.4','卧床、术后、康复期导引',''),('  6.5','安全、强度、时长控制',''),
('第7章','颈肩腰专项康复功法','8课时'),('  7.1','颈椎放松、矫正、养护功法',''),
('  7.2','腰椎强化、核心、护腰功法',''),('  7.3','肩关节松解、活动、防粘连',''),
('  7.4','膝关节养护、肌力、稳关节',''),('  7.5','旧伤、劳损、疼痛调理功法',''),
('第8章','睡眠、减压、安神专项功法','8课时'),('  8.1','睡前放松功：呼吸+拉伸+冥想',''),
('  8.2','焦虑、紧张、烦躁快速安定法',''),('  8.3','心神不宁、多梦、易醒调理',''),
('  8.4','职场5分钟快速减压功',''),('  8.5','退役军人情绪稳定专项功法',''),
('第9章','不同人群功法适配教学','8课时'),('  9.1','老年人功法：温和、慢、少、安全',''),
('  9.2','职工功法：短、快、便、易坚持',''),('  9.3','退役军人功法：强筋骨、祛寒湿、稳状态',''),
('  9.4','慢病患者功法：禁忌、强度、监护',''),('  9.5','青少年功法：助长、护眼、正体态',''),
('第10章','团体带教与活动组织','8课时'),('  10.1','口令设计、示范站位、队形安排',''),
('  10.2','晨练、课堂、沙龙、展演组织',''),('  10.3','安全巡查、风险预警、应急处理',''),
('  10.4','打卡、督导、社群运营',''),('  10.5','师资考核、结业标准、等级评定',''),
]
for num,title,hours in toc:
    p = doc.add_paragraph()
    r = p.add_run(f'{num}  {title}'); r.font.size = Pt(12)
    r.font.name = '宋体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if not num.startswith(' '): r.bold = True
    if hours:
        r2 = p.add_run(f'  …………  {hours}'); r2.font.size = Pt(12)
        r2.font.name = '宋体'; r2.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
doc.add_page_break()

doc.save(OUTPATH)
with open(STATE,'wb') as f: pickle.dump(total_chars, f)
print(f"Base OK, chars={total_chars}")
